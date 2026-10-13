#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Red de pluviómetros de Castilla para Rio Map (ADR-0037), desde los archivos de
Recursos Hídricos de Castilla:

- "Caracteristicas de pluviometros.xlsx": código, zona, nombre, técnico, x, y.
- "Pluviometro suerte.xlsx" (hoja "Suerte-Pluviometro"): suerte → pluviómetro.

Salidas (public/data/):
- pluviometros_castilla.json  [{ id, zona, tecnico, hacienda, sitio, area_ha, lat, lon }]
  (mismo formato que pluviometros_riopaila.json). `hacienda` = nombre del
  pluviómetro; `area_ha` = área neta (maestro) de las suertes asignadas.
- pluviometros_suertes_castilla.json  { "2101-010": 2135, ... }
  Solo suertes con un pluviómetro de la red; las demás las asigna la app al
  pluviómetro más cercano.

Uso:  python scripts/convertir_pluviometros_castilla.py [carpeta]
"""

import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

CARPETA = Path(
    sys.argv[1]
    if len(sys.argv) > 1
    else Path.home()
    / "Documents"
    / "JUNTA DIRECTIVA"
    / "CASTILLA"
    / "Recursos Hidricos"
)
DATA = Path(__file__).resolve().parent.parent / "public" / "data"
SALIDA_RED = DATA / "pluviometros_castilla.json"
SALIDA_SUERTES = DATA / "pluviometros_suertes_castilla.json"
# Caja amplia de la cartografía de Castilla, para descartar coordenadas erradas.
BBOX = (-76.55, 3.0, -76.18, 3.5)


def limpiar(v):
    s = " ".join(("" if v is None else str(v)).split())
    return s or None


def main() -> None:
    wb = openpyxl.load_workbook(
        CARPETA / "Caracteristicas de pluviometros.xlsx",
        read_only=True,
        data_only=True,
    )
    red = []
    for i, row in enumerate(wb.active.iter_rows(values_only=True)):
        if i == 0 or row[0] is None:
            continue
        pid, zona, _depto, nombre, tecnico, x, y = row[:7]
        lon, lat = float(x), float(y)
        if not (BBOX[0] <= lon <= BBOX[2] and BBOX[1] <= lat <= BBOX[3]):
            raise SystemExit(f"PV {pid}: coordenada fuera de Castilla ({lat}, {lon})")
        red.append(
            {
                "id": int(pid),
                "zona": int(zona),
                "tecnico": limpiar(tecnico),
                "hacienda": limpiar(nombre),
                "sitio": None,
                "area_ha": None,
                "lat": round(lat, 7),
                "lon": round(lon, 7),
            }
        )
    ids = {p["id"] for p in red}

    wb = openpyxl.load_workbook(
        CARPETA / "Pluviometro suerte.xlsx", read_only=True, data_only=True
    )
    asignacion = {}
    fuera_de_red = defaultdict(list)
    sin_pluv = []
    for i, row in enumerate(wb["Suerte-Pluviometro"].iter_rows(values_only=True)):
        if i == 0 or not row[0]:
            continue
        suerte = str(row[0]).strip()
        pv = row[1]
        if pv is None:
            sin_pluv.append(suerte)
        elif int(pv) in ids:
            asignacion[suerte] = int(pv)
        else:
            fuera_de_red[int(pv)].append(suerte)

    # Área de influencia: área neta del maestro de las suertes asignadas.
    maestro = json.loads((DATA / "maestro_castilla.json").read_text(encoding="utf-8"))
    area = defaultdict(float)
    for suerte, pv in asignacion.items():
        a = (maestro.get(suerte) or {}).get("area_neta_ha")
        if isinstance(a, (int, float)):
            area[pv] += a
    for p in red:
        p["area_ha"] = round(area[p["id"]], 2) if area[p["id"]] else None

    red.sort(key=lambda r: r["id"])
    SALIDA_RED.write_text(json.dumps(red, ensure_ascii=False, indent=0), encoding="utf-8")
    SALIDA_SUERTES.write_text(
        json.dumps(dict(sorted(asignacion.items())), ensure_ascii=False, indent=0),
        encoding="utf-8",
    )

    mapa = {
        f["properties"]["sec_ste"]
        for f in json.loads(
            (DATA / "tablones_castilla.geojson").read_text(encoding="utf-8")
        )["features"]
    }
    print(f"{SALIDA_RED.name}: {len(red)} pluviómetros")
    print(f"Técnicos: {sorted({p['tecnico'] for p in red})}")
    print(f"{SALIDA_SUERTES.name}: {len(asignacion)} suertes asignadas")
    print(f"  del mapa con asignación oficial: {len(mapa & set(asignacion))} de {len(mapa)}")
    print(f"  sin pluviómetro en el Excel: {len(sin_pluv)}")
    for pv, ss in sorted(fuera_de_red.items()):
        print(f"  PV {pv} no está en la red: {ss}")
    repetidas = defaultdict(list)
    for p in red:
        repetidas[(p["lat"], p["lon"])].append(p["id"])
    for coord, pvs in repetidas.items():
        if len(pvs) > 1:
            print(f"  AVISO: misma coordenada {coord}: PV {pvs}")
    sin_area = [p["id"] for p in red if not p["area_ha"]]
    if sin_area:
        print(f"  PV sin suertes asignadas (sin área): {sin_area}")


if __name__ == "__main__":
    main()
